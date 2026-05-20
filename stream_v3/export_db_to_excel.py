"""
车辆清洗记录数据库导出工具
支持导出为Excel表格，包含格式化、筛选、统计等功能
"""

import sqlite3
import argparse
from datetime import datetime, timedelta
from pathlib import Path
from typing import List, Dict, Optional

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.chart import BarChart, Reference
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False
    print("警告: openpyxl未安装，无法导出Excel")
    print("请运行: pip install openpyxl")


class VehicleWashDBExporter:
    """车辆清洗记录数据库导出器"""

    def __init__(self, db_path: str = 'vehicle_wash.db'):
        self.db_path = db_path
        self.connection = None

    def _connect(self):
        """连接数据库"""
        try:
            self.connection = sqlite3.connect(self.db_path)
            self.connection.row_factory = sqlite3.Row
            print(f"数据库连接成功: {self.db_path}")
        except Exception as e:
            raise Exception(f"数据库连接失败: {e}")

    def _get_records(self, start_date: Optional[str] = None,
                     end_date: Optional[str] = None,
                     plate_filter: Optional[str] = None) -> List[Dict]:
        """获取记录，支持日期和车牌筛选"""
        if not self.connection:
            self._connect()

        sql = "SELECT * FROM vehicle_wash_records WHERE 1=1"
        params = []

        if start_date:
            sql += " AND entry_time >= ?"
            params.append(start_date)
        if end_date:
            sql += " AND entry_time <= ?"
            params.append(end_date)
        if plate_filter:
            sql += " AND license_plate LIKE ?"
            params.append(f"%{plate_filter}%")

        sql += " ORDER BY entry_time DESC"

        try:
            cursor = self.connection.execute(sql, params)
            return [dict(row) for row in cursor.fetchall()]
        except Exception as e:
            print(f"查询记录失败: {e}")
            return []

    def _format_record(self, record: Dict) -> Dict:
        """格式化单条记录，便于展示"""
        formatted = dict(record)

        # 转换时间格式
        for time_field in ['entry_time', 'exit_time', 'wash_start_time', 'created_at', 'updated_at']:
            if formatted.get(time_field):
                try:
                    dt = datetime.fromisoformat(formatted[time_field])
                    formatted[time_field] = dt.strftime('%Y-%m-%d %H:%M:%S')
                except:
                    pass

        # 转换布尔值
        formatted['is_washed'] = '是' if formatted.get('is_washed') == 1 else '否'
        formatted['is_reentry'] = '是' if formatted.get('is_reentry') == 1 else '否'

        # 格式化停留时间
        dwell_time = formatted.get('dwell_time', 0) or 0
        if dwell_time > 0:
            minutes = int(dwell_time // 60)
            seconds = int(dwell_time % 60)
            formatted['dwell_time_formatted'] = f"{minutes}分{seconds}秒"
        else:
            formatted['dwell_time_formatted'] = '-'

        # 格式化清洗时长
        wash_duration = formatted.get('wash_duration', 0) or 0
        if wash_duration > 0:
            minutes = int(wash_duration // 60)
            seconds = int(wash_duration % 60)
            formatted['wash_duration_formatted'] = f"{minutes}分{seconds}秒"
        else:
            formatted['wash_duration_formatted'] = '-'

        return formatted

    def export_to_excel(self, output_path: Optional[str] = None,
                        start_date: Optional[str] = None,
                        end_date: Optional[str] = None,
                        plate_filter: Optional[str] = None) -> str:
        """
        导出记录到Excel

        Args:
            output_path: 输出文件路径，默认为 车辆清洗记录_YYYYMMDD.xlsx
            start_date: 开始日期筛选 (格式: YYYY-MM-DD)
            end_date: 结束日期筛选 (格式: YYYY-MM-DD)
            plate_filter: 车牌号筛选（支持模糊匹配）

        Returns:
            导出的文件路径
        """
        if not EXCEL_AVAILABLE:
            raise ImportError("openpyxl未安装，无法导出Excel")

        # 获取记录
        records = self._get_records(start_date, end_date, plate_filter)
        if not records:
            print("没有找到记录")
            return ""

        print(f"找到 {len(records)} 条记录")

        # 格式化记录
        formatted_records = [self._format_record(r) for r in records]

        # 创建Excel工作簿
        wb = Workbook()
        ws = wb.active
        ws.title = "清洗记录"

        # 定义样式
        header_font = Font(name='微软雅黑', size=11, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

        data_font = Font(name='微软雅黑', size=10)
        data_alignment = Alignment(horizontal='center', vertical='center')
        left_alignment = Alignment(horizontal='left', vertical='center')

        thin_border = Border(
            left=Side(style='thin', color='CCCCCC'),
            right=Side(style='thin', color='CCCCCC'),
            top=Side(style='thin', color='CCCCCC'),
            bottom=Side(style='thin', color='CCCCCC')
        )

        washed_fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
        washed_font = Font(name='微软雅黑', size=10, color='006100')

        # 标题行
        headers = [
            ('序号', 8),
            ('车牌号码', 15),
            ('进场时间', 20),
            ('出场时间', 20),
            ('停留时间', 15),
            ('是否清洗', 12),
            ('清洗开始时间', 20),
            ('清洗时长', 15),
            ('是否二次进场', 14),
            ('车辆ID', 12),
            ('数据库ID', 10),
        ]

        # 写入标题
        for col_idx, (header, width) in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        # 写入数据
        for row_idx, record in enumerate(formatted_records, 2):
            row_data = [
                row_idx - 1,  # 序号
                record.get('license_plate', ''),
                record.get('entry_time', ''),
                record.get('exit_time', '-'),
                record.get('dwell_time_formatted', '-'),
                record.get('is_washed', '否'),
                record.get('wash_start_time', '-'),
                record.get('wash_duration_formatted', '-'),
                record.get('is_reentry', '否'),
                record.get('vehicle_id', ''),
                record.get('id', ''),
            ]

            for col_idx, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.font = data_font
                cell.border = thin_border

                # 设置对齐方式
                if col_idx == 2:  # 车牌号左对齐
                    cell.alignment = left_alignment
                else:
                    cell.alignment = data_alignment

                # 已清洗的行用绿色标记
                if record.get('is_washed') == '是':
                    cell.fill = washed_fill
                    cell.font = washed_font

        # 冻结首行
        ws.freeze_panes = 'A2'

        # 添加筛选
        ws.auto_filter.ref = ws.dimensions

        # 添加统计信息工作表
        self._add_statistics_sheet(wb, formatted_records)

        # 保存文件
        if not output_path:
            today = datetime.now().strftime('%Y%m%d')
            output_path = f"车辆清洗记录_{today}.xlsx"

        wb.save(output_path)
        print(f"导出成功: {output_path}")
        print(f"共导出 {len(records)} 条记录")

        return output_path

    def _add_statistics_sheet(self, wb: Workbook, records: List[Dict]):
        """添加统计信息工作表"""
        ws = wb.create_sheet(title="统计汇总")

        # 计算统计数据
        total = len(records)
        active = sum(1 for r in records if not r.get('exit_time') or r.get('exit_time') == '-')
        exited = total - active
        washed = sum(1 for r in records if r.get('is_washed') == '是')
        reentry = sum(1 for r in records if r.get('is_reentry') == '是')

        # 按日期统计
        date_stats = {}
        for r in records:
            entry_time = r.get('entry_time', '')
            if entry_time:
                date = entry_time[:10]  # YYYY-MM-DD
                if date not in date_stats:
                    date_stats[date] = {'total': 0, 'washed': 0, 'exited': 0}
                date_stats[date]['total'] += 1
                if r.get('is_washed') == '是':
                    date_stats[date]['washed'] += 1
                if r.get('exit_time') and r.get('exit_time') != '-':
                    date_stats[date]['exited'] += 1

        # 样式
        title_font = Font(name='微软雅黑', size=14, bold=True, color='FFFFFF')
        title_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        title_alignment = Alignment(horizontal='center', vertical='center')

        label_font = Font(name='微软雅黑', size=11, bold=True)
        value_font = Font(name='微软雅黑', size=11)

        # 标题
        ws.merge_cells('A1:D1')
        ws['A1'] = '车辆清洗记录统计汇总'
        ws['A1'].font = title_font
        ws['A1'].fill = title_fill
        ws['A1'].alignment = title_alignment
        ws.row_dimensions[1].height = 30

        # 总体统计
        stats_data = [
            ('总记录数', total),
            ('已出场', exited),
            ('在场中', active),
            ('已清洗', washed),
            ('清洗率', f"{washed/exited*100:.1f}%" if exited > 0 else '0%'),
            ('二次进场', reentry),
        ]

        row = 3
        for label, value in stats_data:
            ws.cell(row=row, column=1, value=label).font = label_font
            ws.cell(row=row, column=2, value=value).font = value_font
            row += 1

        # 按日期统计
        row += 1
        ws.cell(row=row, column=1, value='日期').font = label_font
        ws.cell(row=row, column=2, value='进场数').font = label_font
        ws.cell(row=row, column=3, value='出场数').font = label_font
        ws.cell(row=row, column=4, value='清洗数').font = label_font
        ws.cell(row=row, column=5, value='清洗率').font = label_font
        row += 1

        for date in sorted(date_stats.keys(), reverse=True):
            stats = date_stats[date]
            ws.cell(row=row, column=1, value=date)
            ws.cell(row=row, column=2, value=stats['total'])
            ws.cell(row=row, column=3, value=stats['exited'])
            ws.cell(row=row, column=4, value=stats['washed'])
            rate = f"{stats['washed']/stats['exited']*100:.1f}%" if stats['exited'] > 0 else '0%'
            ws.cell(row=row, column=5, value=rate)
            row += 1

        # 设置列宽
        ws.column_dimensions['A'].width = 20
        ws.column_dimensions['B'].width = 15
        ws.column_dimensions['C'].width = 15
        ws.column_dimensions['D'].width = 15
        ws.column_dimensions['E'].width = 15

    def close(self):
        """关闭数据库连接"""
        if self.connection:
            self.connection.close()
            print("数据库连接已关闭")


def main():
    """命令行入口"""
    parser = argparse.ArgumentParser(description='车辆清洗记录数据库导出工具')
    parser.add_argument('--db', default='vehicle_wash.db', help='数据库文件路径')
    parser.add_argument('--output', '-o', help='输出Excel文件路径')
    parser.add_argument('--start-date', help='开始日期 (YYYY-MM-DD)')
    parser.add_argument('--end-date', help='结束日期 (YYYY-MM-DD)')
    parser.add_argument('--plate', help='车牌号筛选（支持模糊匹配）')
    parser.add_argument('--today', action='store_true', help='只导出今天的记录')
    parser.add_argument('--yesterday', action='store_true', help='只导出昨天的记录')
    parser.add_argument('--this-week', action='store_true', help='只导出本周的记录')

    args = parser.parse_args()

    # 处理快捷日期选项
    start_date = args.start_date
    end_date = args.end_date

    if args.today:
        today = datetime.now().strftime('%Y-%m-%d')
        start_date = f"{today} 00:00:00"
        end_date = f"{today} 23:59:59"
    elif args.yesterday:
        yesterday = (datetime.now() - timedelta(days=1)).strftime('%Y-%m-%d')
        start_date = f"{yesterday} 00:00:00"
        end_date = f"{yesterday} 23:59:59"
    elif args.this_week:
        today = datetime.now()
        monday = today - timedelta(days=today.weekday())
        start_date = monday.strftime('%Y-%m-%d 00:00:00')
        end_date = today.strftime('%Y-%m-%d 23:59:59')

    # 导出
    exporter = VehicleWashDBExporter(args.db)
    try:
        output_path = exporter.export_to_excel(
            output_path=args.output,
            start_date=start_date,
            end_date=end_date,
            plate_filter=args.plate
        )
        if output_path:
            print(f"\n导出文件: {output_path}")
    finally:
        exporter.close()


if __name__ == '__main__':
    main()
