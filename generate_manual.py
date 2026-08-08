# -*- coding: utf-8 -*-
"""
生成《车辆清洗智能识别系统产品说明书》Word 文档
面向非技术操作人员，避免使用专业技术术语
"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUTPUT_PATH = r"c:\Users\Admini503\.openclaw\workspace\车辆清洗智能识别系统产品说明书.docx"


def set_cell_shading(cell, fill):
    """设置单元格背景色"""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def set_cell_border(cell, **kwargs):
    """设置单元格边框"""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = 'w:{}'.format(edge)
            element = OxmlElement(tag)
            element.set(qn('w:val'), edge_data.get('val', 'single'))
            element.set(qn('w:sz'), str(edge_data.get('sz', 4)))
            element.set(qn('w:space'), '0')
            element.set(qn('w:color'), edge_data.get('color', '000000'))
            tcBorders.append(element)
    tcPr.append(tcBorders)


def add_heading_custom(doc, text, level):
    """添加统一格式标题"""
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.name = 'Microsoft YaHei'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        if level == 1:
            run.font.size = Pt(18)
            run.font.color.rgb = RGBColor(0, 51, 102)
            run.font.bold = True
        elif level == 2:
            run.font.size = Pt(14)
            run.font.color.rgb = RGBColor(0, 102, 153)
            run.font.bold = True
        else:
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(0, 0, 0)
            run.font.bold = True
    heading.paragraph_format.space_before = Pt(16)
    heading.paragraph_format.space_after = Pt(8)
    return heading


def add_paragraph_custom(doc, text, bold=False, indent=True, align=WD_ALIGN_PARAGRAPH.LEFT):
    """添加统一格式正文段落"""
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_after = Pt(6)
    p.alignment = align
    run = p.add_run(text)
    run.font.name = 'Microsoft YaHei'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    run.font.size = Pt(11)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p


def add_bullet(doc, text, level=0):
    """添加项目符号列表"""
    p = doc.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet 2')
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Cm(0.74 + level * 0.5)
    run = p.add_run(text)
    run.font.name = 'Microsoft YaHei'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p


def add_numbered(doc, text):
    """添加编号列表"""
    p = doc.add_paragraph(style='List Number')
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = 'Microsoft YaHei'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p


def add_table_custom(doc, rows, cols, col_widths=None):
    """添加统一格式表格"""
    table = doc.add_table(rows=rows, cols=cols)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    if col_widths:
        for i, width in enumerate(col_widths):
            for cell in table.columns[i].cells:
                cell.width = Cm(width)
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.space_before = Pt(3)
                for run in paragraph.runs:
                    run.font.name = 'Microsoft YaHei'
                    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
                    run.font.size = Pt(10.5)
    return table


def fill_table_header(table, headers):
    """填充表头并设置样式"""
    hdr_cells = table.rows[0].cells
    for i, text in enumerate(headers):
        hdr_cells[i].text = text
        set_cell_shading(hdr_cells[i], 'E7E6E6')
        for paragraph in hdr_cells[i].paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.bold = True
                run.font.size = Pt(10.5)


def fill_table_row(table, row_idx, values):
    """填充表格行"""
    cells = table.rows[row_idx].cells
    for i, text in enumerate(values):
        cells[i].text = str(text)
        cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_document_default_font(doc):
    """设置文档默认字体"""
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Microsoft YaHei'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0, 0, 0)
    style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')


def add_page_break(doc):
    doc.add_page_break()


def main():
    doc = Document()
    set_document_default_font(doc)

    # 页面边距
    sections = doc.sections
    for section in sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.17)
        section.right_margin = Cm(3.17)

    # ==================== 封面 ====================
    for _ in range(6):
        doc.add_paragraph()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run('车辆清洗智能识别系统')
    run.font.name = 'Microsoft YaHei'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    run.font.size = Pt(32)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 51, 102)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run('产品说明书')
    run.font.name = 'Microsoft YaHei'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(0, 102, 153)

    doc.add_paragraph()
    doc.add_paragraph()

    info_table = doc.add_table(rows=4, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_table.autofit = False
    for row in info_table.rows:
        row.height = Cm(0.8)
    info_data = [
        ('适用场景', '建筑工地、渣土车场、矿山出入口等车辆清洗监管场景'),
        ('前端设备', '500 万像素高清摄像机 + 边缘智能分析主机'),
        ('部署方式', '端侧部署，各项目独立运行、独立记录'),
        ('文档版本', 'V3.0'),
    ]
    for i, (k, v) in enumerate(info_data):
        cells = info_table.rows[i].cells
        cells[0].text = k
        cells[1].text = v
        cells[0].width = Cm(3.5)
        cells[1].width = Cm(10)
        set_cell_shading(cells[0], 'F2F2F2')
        for paragraph in cells[0].paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.bold = True
                run.font.name = 'Microsoft YaHei'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
                run.font.size = Pt(11)
        for paragraph in cells[1].paragraphs:
            for run in paragraph.runs:
                run.font.name = 'Microsoft YaHei'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
                run.font.size = Pt(11)

    doc.add_paragraph()
    doc.add_paragraph()
    date_p = doc.add_paragraph()
    date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = date_p.add_run('2026 年 8 月')
    run.font.name = 'Microsoft YaHei'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(80, 80, 80)

    add_page_break(doc)

    # ==================== 目录占位 ====================
    add_heading_custom(doc, '目录', 1)
    toc_items = [
        '1. 产品概述',
        '2. 系统组成',
        '3. 功能说明',
        '   3.1 视频接入与实时分析',
        '   3.2 车辆识别',
        '   3.3 进出场管理',
        '   3.4 清洗判定',
        '   3.5 数据上报',
        '   3.6 视频回传',
        '   3.7 本地记录与查询',
        '   3.8 多项目管理',
        '4. 安装部署',
        '5. 使用说明',
        '6. 日常维护',
        '7. 常见问题',
        '8. 技术参数',
        '附录：术语速查表',
    ]
    for item in toc_items:
        p = doc.add_paragraph(item)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        p.paragraph_format.space_after = Pt(3)
        for run in p.runs:
            run.font.name = 'Microsoft YaHei'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
            run.font.size = Pt(11)

    add_page_break(doc)

    # ==================== 1. 产品概述 ====================
    add_heading_custom(doc, '1. 产品概述', 1)
    add_paragraph_custom(doc,
        '车辆清洗智能识别系统是一款部署在工地、渣土车场等场所出入口的智能监管设备。'
        '系统通过高清摄像机实时拍摄车辆进出画面，自动识别车辆号牌，判断车辆是否经过清洗区域并停留足够时间，'
        '最终生成车辆进出场记录和清洗记录，并可将数据上报至上级管理平台。')

    add_heading_custom(doc, '1.1 产品定位', 2)
    add_paragraph_custom(doc,
        '本产品定位于施工现场车辆清洗管理的智能化改造。传统人工登记方式存在漏记、错记、难以追溯等问题，'
        '本系统通过自动识别和记录，帮助管理人员实时掌握车辆清洗情况，提升环保合规管理水平。')

    add_heading_custom(doc, '1.2 应用场景', 2)
    add_bullet(doc, '建筑工地出入口车辆清洗监管')
    add_bullet(doc, '渣土运输车辆出场清洗判定')
    add_bullet(doc, '矿山、堆场等重载车辆清洗管理')
    add_bullet(doc, '需要自动记录车辆进出场和清洗状态的场所')

    add_heading_custom(doc, '1.3 核心价值', 2)
    add_bullet(doc, '自动识别：无需人工干预，自动识别车辆号牌和清洗行为')
    add_bullet(doc, '实时监管：车辆进场、出场、清洗状态实时更新')
    add_bullet(doc, '有据可查：每条记录均带抓拍图片和时间戳，支持事后追溯')
    add_bullet(doc, '多端上报：支持同时向多个管理平台报送数据')
    add_bullet(doc, '独立运行：每个项目配置独立的数据和日志，互不干扰')

    add_page_break(doc)

    # ==================== 2. 系统组成 ====================
    add_heading_custom(doc, '2. 系统组成', 1)
    add_paragraph_custom(doc,
        '系统由前端感知设备、边缘智能分析主机和管理平台三部分组成。各组成部分协同工作，'
        '完成视频采集、智能分析、数据存储和数据上报等全部功能。')

    add_heading_custom(doc, '2.1 前端感知设备', 2)
    add_paragraph_custom(doc,
        '前端感知设备为 500 万像素高清摄像机，安装于清洗区域或出入口上方，用于全天候拍摄车辆进出画面。'
        '摄像机支持通过网络实时传输视频画面，保证分析主机能够持续获取清晰、稳定的图像。')

    add_heading_custom(doc, '2.2 边缘智能分析主机', 2)
    add_paragraph_custom(doc,
        '边缘智能分析主机部署在施工现场本地，负责接收摄像机画面，运行车辆识别、车牌识别、进出场判定、'
        '清洗判定等智能分析算法，并将分析结果保存到本地，同时按需上报管理平台。'
        '主机运行稳定，可在工业现场长期无人值守工作。')

    add_heading_custom(doc, '2.3 管理平台', 2)
    add_paragraph_custom(doc,
        '管理平台用于接收、展示和统计各项目上报的车辆进出场及清洗数据。系统支持同时向两个管理平台报送数据，'
        '满足不同层级监管需求。管理平台可通过电脑浏览器访问，实现远程监管。')

    comp_table = add_table_custom(doc, 4, 3, col_widths=[3.5, 5.5, 6])
    fill_table_header(comp_table, ['组成部分', '部署位置', '主要作用'])
    fill_table_row(comp_table, 1, ['高清摄像机', '现场清洗区/出入口', '实时采集车辆视频画面'])
    fill_table_row(comp_table, 2, ['边缘智能分析主机', '现场机房/配电间', '运行智能分析、保存记录、上报数据'])
    fill_table_row(comp_table, 3, ['管理平台', '监管中心/云端', '接收数据、展示统计、远程监管'])

    add_page_break(doc)

    # ==================== 3. 功能说明 ====================
    add_heading_custom(doc, '3. 功能说明', 1)
    add_paragraph_custom(doc,
        '本章详细说明系统的主要功能。各项功能均围绕“车辆识别—进出场判定—清洗判定—数据上报”这一核心业务流程展开。')

    # 3.1
    add_heading_custom(doc, '3.1 视频接入与实时分析', 2)
    add_paragraph_custom(doc,
        '系统支持接入符合行业通用标准的网络视频流。摄像机将现场画面实时传输至边缘智能分析主机，'
        '主机对每一帧画面进行连续分析。')
    add_bullet(doc, '支持接入常见网络摄像机或流媒体视频源')
    add_bullet(doc, '当网络出现短暂中断时，系统会自动尝试重新连接视频源，恢复正常后自动继续分析')
    add_bullet(doc, '支持本地视频文件回放分析，用于调试和复核')

    # 3.2
    add_heading_custom(doc, '3.2 车辆识别', 2)
    add_paragraph_custom(doc,
        '系统内置车辆识别算法，能够自动从画面中识别出进入视野的车辆，并进一步识别车辆号牌信息。')
    add_bullet(doc, '可识别小汽车、货车、客车、摩托车等多种车辆类型')
    add_bullet(doc, '自动识别车辆号牌号码和号牌颜色')
    add_bullet(doc, '对号牌识别结果进行校验，过滤模糊、错误识别结果')
    add_bullet(doc, '同一车辆在多帧画面中被连续识别时，系统会自动合并为一次完整记录，避免重复上报')

    add_paragraph_custom(doc,
        '车辆识别过程在本地完成，识别结果仅在必要时上传图片和文字信息，既保障了数据安全，也降低了对网络带宽的占用。')

    # 3.3
    add_heading_custom(doc, '3.3 进出场管理', 2)
    add_paragraph_custom(doc,
        '系统根据车辆在画面中的位置和移动方向，自动判定车辆进场和出场，并记录相应时间。')

    add_heading_custom(doc, '3.3.1 进场判定', 3)
    add_paragraph_custom(doc,
        '当车辆驶入监控画面，并且车牌被稳定识别后，系统即判定该车辆“进场”，同时保存进场时间和进场抓拍图片。'
        '若同一车牌在出场后再次进入画面，系统会识别为“二次进场”，并重新建立记录。')

    add_heading_custom(doc, '3.3.2 出场判定', 3)
    add_paragraph_custom(doc,
        '系统根据项目实际情况配置车辆离场方向，可设置为从画面左侧、右侧或下方离场。'
        '当车辆前缘越过预设的出口线并持续一定时间后，系统判定车辆“出场”。'
        '为避免误判，系统还设置了最短在场时间校验，确保车辆在场内停留足够时间后才允许判定出场。')

    add_heading_custom(doc, '3.3.3 停留时间统计', 3)
    add_paragraph_custom(doc,
        '车辆出场后，系统自动计算从进场到出场的总停留时间，并写入记录。管理人员可通过管理平台查看每辆车的停留时长。')

    # 3.4
    add_heading_custom(doc, '3.4 清洗判定', 2)
    add_paragraph_custom(doc,
        '清洗判定是本系统的核心功能之一。系统在画面中划定清洗区域，当车辆进入该区域并停止足够时间，'
        '即判定车辆完成清洗。')

    add_bullet(doc, '清洗区域可在项目配置中设定，适应不同现场布局')
    add_bullet(doc, '车辆进入清洗区域并停止后，系统开始计时')
    add_bullet(doc, '当车辆在区域内持续停止达到设定时长，系统判定为“已清洗”')
    add_bullet(doc, '若车辆在计时期间离开清洗区域或明显移动，计时会自动重置')

    add_paragraph_custom(doc,
        '清洗判定完成后，系统会在车辆出场时将“是否清洗”状态随出场记录一同上报。')

    # 3.5
    add_heading_custom(doc, '3.5 数据上报', 2)
    add_paragraph_custom(doc,
        '系统可将车辆进出场记录和清洗状态上报至管理平台。每条上报记录均包含设备编号、车牌号码、进出场类型、'
        '是否清洗、抓拍图片和发生时间等信息。')

    add_bullet(doc, '进场时上报：车辆被确认进场后，自动上报进场记录和抓拍图片')
    add_bullet(doc, '出场时上报：车辆判定出场后，自动上报出场记录、停留时间和清洗状态')
    add_bullet(doc, '双平台上报：支持同时向两个管理平台报送数据，满足多级监管需求')
    add_bullet(doc, '上报失败自动重试：网络异常时系统会记录失败次数，待网络恢复后继续尝试')

    add_paragraph_custom(doc,
        '数据上报过程对操作人员透明，系统启动后自动运行，无需人工干预。')

    # 3.6
    add_heading_custom(doc, '3.6 视频回传', 2)
    add_paragraph_custom(doc,
        '系统可将标注了车辆检测框、车牌、清洗区域和出口线等信息的实时画面回传至视频平台，'
        '方便管理人员远程查看现场情况。')

    add_bullet(doc, '实时画面上叠加车辆框、车牌号码、清洗区域边界和出口线')
    add_bullet(doc, '支持将画面推送到标准流媒体服务器，供电脑、手机等终端观看')
    add_bullet(doc, '推流过程具备自动重连能力，网络波动恢复后可继续推送')
    add_bullet(doc, '画面分辨率、帧率等参数可按现场网络条件配置')

    # 3.7
    add_heading_custom(doc, '3.7 本地记录与查询', 2)
    add_paragraph_custom(doc,
        '系统在每个项目中维护独立的本地记录，即使网络中断，车辆进出场和清洗数据也不会丢失，'
        '待网络恢复后继续上报。')

    add_bullet(doc, '记录内容包括：车牌号、进场时间、出场时间、停留时长、是否清洗、清洗开始时间、清洗时长等')
    add_bullet(doc, '支持查询当前在场车辆')
    add_bullet(doc, '支持查询历史记录和统计信息')
    add_bullet(doc, '各项目数据独立存储，互不干扰')

    # 3.8
    add_heading_custom(doc, '3.8 多项目管理', 2)
    add_paragraph_custom(doc,
        '一台边缘智能分析主机可同时管理多个项目，每个项目使用独立的配置文件、独立日志和独立数据记录。'
        '启动后，各项目并行运行，适用于集团化、多工地集中管理的场景。')

    add_bullet(doc, '每个项目可配置独立的摄像机、清洗区域、出口方向和上报平台')
    add_bullet(doc, '项目之间数据隔离，便于分项目统计和追溯')
    add_bullet(doc, '可通过统一启动器批量启停所有项目')

    add_page_break(doc)

    # ==================== 4. 安装部署 ====================
    add_heading_custom(doc, '4. 安装部署', 1)

    add_heading_custom(doc, '4.1 部署要求', 2)
    req_table = add_table_custom(doc, 5, 2, col_widths=[4, 10])
    fill_table_header(req_table, ['项目', '要求'])
    fill_table_row(req_table, 1, ['边缘智能分析主机', '运行稳定的嵌入式工业主机，建议配备图形运算加速模块'])
    fill_table_row(req_table, 2, ['高清摄像机', '500 万像素及以上，网络型摄像机，覆盖清洗区和出入口'])
    fill_table_row(req_table, 3, ['网络', '摄像机与分析主机之间网络通畅；分析主机可访问上报平台'])
    fill_table_row(req_table, 4, ['电源', '稳定供电，建议配置不间断电源，防止意外断电导致数据丢失'])

    add_heading_custom(doc, '4.2 摄像机安装建议', 2)
    add_bullet(doc, '安装高度：建议 3~6 米，根据现场车道宽度和车辆高度调整')
    add_bullet(doc, '安装位置：位于车辆行驶方向正前方或侧上方，确保车牌清晰可见')
    add_bullet(doc, '视野范围：画面应完整覆盖车辆进出口和清洗区域')
    add_bullet(doc, '光照条件：避免强光直射和逆光，夜间需有补光设施')
    add_bullet(doc, '网络连接：摄像机与分析主机处于同一局域网或可通过互联网稳定访问')

    add_heading_custom(doc, '4.3 主机安装建议', 2)
    add_bullet(doc, '放置在通风、干燥的机房或配电箱内')
    add_bullet(doc, '确保主机与摄像机之间网络畅通')
    add_bullet(doc, '为主机配置固定 IP 地址，便于远程维护')
    add_bullet(doc, '定期清理主机散热口灰尘，保证散热良好')

    add_page_break(doc)

    # ==================== 5. 使用说明 ====================
    add_heading_custom(doc, '5. 使用说明', 1)

    add_heading_custom(doc, '5.1 启动系统', 2)
    add_numbered(doc, '确认摄像机和分析主机已通电并连接网络。')
    add_numbered(doc, '在分析主机上运行启动程序，系统会自动加载各项目配置并开始工作。')
    add_numbered(doc, '启动后，系统开始接收视频画面并进行分析。')
    add_numbered(doc, '如需停止系统，可发送停止信号，系统会安全关闭各项目进程。')

    add_paragraph_custom(doc,
        '系统启动后会自动完成车辆识别、进出场判定、清洗判定和数据上报，操作人员无需频繁干预。')

    add_heading_custom(doc, '5.2 查看运行状态', 2)
    add_paragraph_custom(doc,
        '系统运行过程中会生成日志文件，记录视频连接、车辆识别、进出场、清洗判定、数据上报和推流等状态。'
        '操作人员可通过查看日志了解系统是否正常运行。')
    add_bullet(doc, '日志文件按项目独立保存，便于定位问题')
    add_bullet(doc, '日志中包含成功识别车辆数、上报成功/失败次数、推流状态等统计信息')
    add_bullet(doc, '如发现持续上报失败，可检查网络连接或管理平台地址配置')

    add_heading_custom(doc, '5.3 查看实时画面', 2)
    add_paragraph_custom(doc,
        '系统支持在本地显示分析画面，也可通过网络视频平台远程查看。'
        '画面上会显示车辆检测框、车牌号码、清洗区域和出口线，方便直观了解系统运行状态。')
    add_bullet(doc, '本地显示窗口可在配置中开启或关闭')
    add_bullet(doc, '多项目同时运行时，建议只保留一个项目的本地显示窗口，减少主机负担')
    add_bullet(doc, '远程查看需通过流媒体服务器地址访问')

    add_heading_custom(doc, '5.4 查看记录', 2)
    add_paragraph_custom(doc,
        '车辆进出场和清洗记录保存在本地，也可通过管理平台上报数据查看。'
        '本地记录可通过配套工具导出为表格，方便进行离线统计和分析。')

    add_page_break(doc)

    # ==================== 6. 日常维护 ====================
    add_heading_custom(doc, '6. 日常维护', 1)

    add_heading_custom(doc, '6.1 日常检查项', 2)
    check_table = add_table_custom(doc, 6, 3, col_widths=[2.5, 7, 4])
    fill_table_header(check_table, ['检查项', '检查内容', '建议周期'])
    fill_table_row(check_table, 1, ['摄像机', '画面是否清晰、有无遮挡、角度是否偏移', '每日'])
    fill_table_row(check_table, 2, ['补光灯', '夜间补光是否正常', '每日'])
    fill_table_row(check_table, 3, ['网络', '摄像机与主机、主机与平台之间网络是否通畅', '每日'])
    fill_table_row(check_table, 4, ['主机', '运行指示灯是否正常、散热是否良好', '每周'])
    fill_table_row(check_table, 5, ['日志', '查看是否有持续报错或上报失败', '每周'])

    add_heading_custom(doc, '6.2 数据备份', 2)
    add_paragraph_custom(doc,
        '建议定期将本地记录文件和日志文件备份到外部存储或管理平台，防止因主机故障导致历史数据丢失。'
        '数据文件可按项目名称查找，备份时注意不要遗漏。')

    add_heading_custom(doc, '6.3 配置更新', 2)
    add_paragraph_custom(doc,
        '当现场布局、摄像机位置、出口方向或上报平台地址发生变化时，需要更新项目配置文件。'
        '配置更新应由技术人员完成，更新后重启对应项目即可生效。')

    add_page_break(doc)

    # ==================== 7. 常见问题 ====================
    add_heading_custom(doc, '7. 常见问题', 1)

    faq_table = add_table_custom(doc, 7, 3, col_widths=[3.5, 5.5, 5])
    fill_table_header(faq_table, ['现象', '可能原因', '处理方法'])
    fill_table_row(faq_table, 1, ['系统无法识别车辆', '摄像机画面模糊、角度偏移或光线不足', '清洁镜头、调整角度、改善光照'])
    fill_table_row(faq_table, 2, ['车牌识别率低', '车牌污损、反光或被遮挡', '清洗车辆号牌、调整摄像机角度避免反光'])
    fill_table_row(faq_table, 3, ['数据上报失败', '网络中断或平台上报地址错误', '检查网络连接、核对平台地址配置'])
    fill_table_row(faq_table, 4, ['清洗判定不准确', '清洗区域设置与实际区域不符', '重新标定清洗区域'])
    fill_table_row(faq_table, 5, ['实时画面无法查看', '推流地址错误或网络不通', '检查推流配置和网络'])
    fill_table_row(faq_table, 6, ['车辆重复进场/出场', '车辆在场内停留时间过短或识别不稳定', '调整最短在场时间和出口确认时间'])

    add_page_break(doc)

    # ==================== 8. 技术参数 ====================
    add_heading_custom(doc, '8. 技术参数', 1)

    param_table = add_table_custom(doc, 13, 2, col_widths=[4, 10])
    fill_table_header(param_table, ['参数项', '参数值/说明'])
    fill_table_row(param_table, 1, ['系统名称', '车辆清洗智能识别系统 V3'])
    fill_table_row(param_table, 2, ['部署方式', '端侧部署，本地智能分析'])
    fill_table_row(param_table, 3, ['前端摄像机', '500 万像素高清网络摄像机'])
    fill_table_row(param_table, 4, ['识别对象', '小汽车、货车、客车、摩托车等'])
    fill_table_row(param_table, 5, ['识别内容', '车辆号牌号码、号牌颜色'])
    fill_table_row(param_table, 6, ['进出场方向', '支持从画面左侧、右侧或下方离场'])
    fill_table_row(param_table, 7, ['清洗判定', '车辆在清洗区域内停止达到设定时长'])
    fill_table_row(param_table, 8, ['数据上报', '支持同时向两个管理平台报送'])
    fill_table_row(param_table, 9, ['视频回传', '支持实时标注画面回传'])
    fill_table_row(param_table, 10, ['本地记录', '按项目独立保存，支持查询和导出'])
    fill_table_row(param_table, 11, ['多项目支持', '单台主机可同时运行多个项目'])
    fill_table_row(param_table, 12, ['异常恢复', '视频断流和推流断开后自动重连'])

    add_page_break(doc)

    # ==================== 附录 ====================
    add_heading_custom(doc, '附录：术语速查表', 1)
    add_paragraph_custom(doc,
        '本附录对说明书中出现的部分名词做简要解释，便于非技术人员理解。')

    term_table = add_table_custom(doc, 9, 2, col_widths=[3.5, 10.5])
    fill_table_header(term_table, ['术语', '说明'])
    fill_table_row(term_table, 1, ['边缘智能分析主机', '部署在现场的本地计算设备，负责运行识别和分析程序'])
    fill_table_row(term_table, 2, ['高清摄像机', '用于拍摄现场画面的网络摄像机，本系统采用 500 万像素机型'])
    fill_table_row(term_table, 3, ['清洗区域', '画面中划定的车辆应当停留清洗的矩形区域'])
    fill_table_row(term_table, 4, ['出口线', '画面中用于判断车辆是否离场的参考线'])
    fill_table_row(term_table, 5, ['进场', '车辆进入监控画面并被识别到号牌'])
    fill_table_row(term_table, 6, ['出场', '车辆前缘越过出口线并持续一定时间后判定为离开'])
    fill_table_row(term_table, 7, ['停留时间', '从进场到出场的总时长'])
    fill_table_row(term_table, 8, ['数据上报', '将车辆进出场和清洗记录发送给管理平台'])

    doc.save(OUTPUT_PATH)
    print(f"产品说明书已生成：{OUTPUT_PATH}")


if __name__ == '__main__':
    main()
